"""Insert LaTeX formulas into python-pptx slides the same way the Latex4PowerPoint add-in does.

The add-in can edit the result in PowerPoint. Run this file directly to test the setup:

    python latex_eq.py
"""
import io
import os
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Emu

L4P = Path(os.environ["APPDATA"]) / "Latex4PowerPoint"  # the add-in's settings.xml and LatexTemplate.txt
SIZE, DPI = 32, 384  # render settings of the user's own formulas; the picture is then scaled to the requested pt
TAGS_CT = "application/vnd.openxmlformats-officedocument.presentationml.tags+xml"


def latex_png(code, color="0,0,0"):
    """Render code with the add-in's template, MiKTeX and settings. Returns PNG bytes."""
    settings, template = L4P / "settings.xml", L4P / "LatexTemplate.txt"
    for f in (settings, template):
        if not f.exists():
            raise FileNotFoundError(f"{f} is missing. Start PowerPoint, open the Latex4PowerPoint dialog once "
                                    "and set the MiKTeX path in its options.")
    miktex = Path(ET.parse(settings).getroot().findtext("miktexPath"))
    tpl = template.read_text(encoding="utf-8")
    body = "\\definecolor{txtcolor}{rgb}{%s}\n\\color{txtcolor}\n\\changefont{cmr}{m}{n}\n%s" % (color, code)
    with tempfile.TemporaryDirectory() as d:
        (Path(d) / "eq.tex").write_text(tpl.replace("${Content}", body), encoding="utf-8")
        for cmd in ([miktex / "latex.exe", "-interaction=nonstopmode", "-halt-on-error", "eq.tex"],
                    [miktex / "dvipng.exe", "-T", "tight", "-bg", "Transparent", "--depth", "--noghostscript",
                     "-D", str(DPI * SIZE // 12), "-o", "eq.png", "eq.dvi"]):
            r = subprocess.run(cmd, cwd=d, capture_output=True, text=True, errors="replace")
            if r.returncode:
                hint = "" if "cgmathnotation" in tpl else (
                    "\nLatexTemplate.txt does not load cgmathnotation: run latex-sty\\install.bat in the repository.")
                raise RuntimeError(f"{cmd[0].name} failed for {code!r}:\n{r.stdout[-1500:]}{hint}")
        return (Path(d) / "eq.png").read_bytes()


def add_latex(slide, code, x, y, pt=20, color="0,0,0"):
    """Add code as a formula picture with its top-left at (x, y) EMU, shown at pt points.

    color: hex like "2F363C" or the add-in's "r,g,b" floats. Returns the picture shape.
    """
    if "," not in color:
        color = color.lstrip("#")
        color = ",".join(f"{int(color[i:i + 2], 16) / 255:.3f}" for i in (0, 2, 4))
    png = latex_png(code, color)
    w, h = Image.open(io.BytesIO(png)).size
    k = 914400 / DPI * pt / SIZE  # EMU per pixel: the add-in shows its PNG at SIZE pt
    pic = slide.shapes.add_picture(io.BytesIO(png), Emu(x), Emu(y), Emu(round(w * k)), Emu(round(h * k)))
    pic._element.spPr.set("bwMode", "black")  # the add-in sets BlackWhiteMode = black
    # The add-in's edit dialog reads all of these tags. Newlines in the code are stored as the 4 characters \r\n.
    tags = dict(LATEXCODE=code.replace("\r\n", "\n").replace("\n", r"\r\n"), LATEXFONTSIZE=str(SIZE),
                LATEXTEXTCOLOR=color, LATEXDPI=str(DPI), LATEXFONT="Computer Modern Roman",
                LATEXFONTSERIES="Standard", LATEXFONTSHAPE="Standard", LATEXTEXTSHAPEID="-1",
                LATEXADDINVERSION="1.3.0.0")
    lst = etree.Element(qn("p:tagLst"), nsmap={"p": "http://schemas.openxmlformats.org/presentationml/2006/main"})
    for name, val in tags.items():
        etree.SubElement(lst, qn("p:tag"), name=name, val=val)
    pkg = slide.part.package
    part = Part(pkg.next_partname("/ppt/tags/tag%d.xml"), TAGS_CT, pkg,
                etree.tostring(lst, xml_declaration=True, encoding="UTF-8", standalone=True))
    cdl = etree.SubElement(pic._element.nvPicPr.nvPr, qn("p:custDataLst"))
    etree.SubElement(cdl, qn("p:tags")).set(qn("r:id"), slide.part.relate_to(part, RT.TAGS))
    return pic


if __name__ == "__main__":  # self-check: one cgmathnotation formula into a temporary deck
    import zipfile

    from pptx import Presentation

    prs = Presentation()
    pic = add_latex(prs.slides.add_slide(prs.slide_layouts[6]), r"$\trans{\mat{F}}\mat{F}=\mat{C}$", 914400, 914400)
    out = Path(tempfile.gettempdir()) / "latex_eq_selfcheck.pptx"
    prs.save(out)
    with zipfile.ZipFile(out) as z:
        tags = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("ppt/tags/")]
    assert len(tags) == 1 and r"\trans{\mat{F}}" in tags[0], tags
    print(f"OK: formula rendered and tagged ({pic.width / 12700:.0f} x {pic.height / 12700:.0f} pt). "
          f"Open {out} in PowerPoint to try the add-in's edit button.")
